"""
Import actual trade history from ASSET ALLOCATION.xlsx (Journal-2 sheet)
into psx_data.db as source='Actual' trades.

Column mapping:
  Date1      → entry date (buy date)
  Name       → symbol / stock ticker
  Status     → C=Closed, O=Open, blank=not initiated (blank rows skipped)
  QTY        → quantity
  BUY        → entry price
  Value      → position value
  SL         → initial stop loss price
  POS        → Long or Short → LONG / SHORT
  Date2      → exit date (sell date)
  SELL       → exit price
  Gain/Loss  → P&L in PKR
  Gain/Loss% → P&L percentage
  Days       → holding duration
  R:R        → actual risk/reward achieved

Usage:
    python import_actual_trades.py --file "C:\\path\\to\\ASSET ALLOCATION.xlsx" --dry-run
    python import_actual_trades.py --file "C:\\path\\to\\ASSET ALLOCATION.xlsx"
"""

import argparse
import logging
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    handlers=[logging.StreamHandler(sys.stdout)],
)
logger = logging.getLogger(__name__)

import pandas as pd
from database import get_conn, init_db

# Guard: this script uses SQLite-only ? placeholders and must NOT run against PostgreSQL.
# Running it with DATABASE_URL set would cause ProgrammingError on every row.
if os.environ.get("DATABASE_URL") or os.environ.get("SUPABASE_DB_URL"):
    sys.exit(
        "ERROR: import_actual_trades.py uses SQLite syntax and cannot run when "
        "DATABASE_URL is set. Unset DATABASE_URL before running this script locally."
    )

SHEET_NAME = "JOURNAL-2"


def _parse_date(val) -> str | None:
    if pd.isna(val) or str(val).strip() in ("", "nan", "NaT"):
        return None
    try:
        return pd.to_datetime(val, dayfirst=True).strftime("%Y-%m-%d")
    except Exception:
        return None


def _normalise_direction(val) -> str:
    v = str(val).strip().lower()
    if v in ("short", "s", "sell"):
        return "SHORT"
    return "LONG"


def _outcome_from_pl(pl_pct) -> str | None:
    try:
        v = float(pl_pct)
        if v > 0.0:
            return "Win"
        if v < 0.0:
            return "Loss"
        return "Breakeven"
    except Exception:
        return None


def import_trades(filepath: str, dry_run: bool = False) -> dict:
    logger.info("Reading sheet '%s' from: %s", SHEET_NAME, filepath)

    try:
        # Read without header first to locate the actual header row
        raw = pd.read_excel(filepath, sheet_name=SHEET_NAME, header=None)
    except Exception as e:
        logger.error("Could not read sheet '%s': %s", SHEET_NAME, e)
        logger.info("Available sheets: %s", pd.ExcelFile(filepath).sheet_names)
        return {"imported": 0, "skipped": 0, "errors": 1}

    # Find the row that contains "Date1" (the column header row)
    header_row = None
    for i, row in raw.iterrows():
        if any(str(cell).strip() == "Date1" for cell in row):
            header_row = i
            break

    if header_row is None:
        logger.error("Could not find header row with 'Date1' in sheet '%s'", SHEET_NAME)
        logger.info("First 5 rows: %s", raw.head().to_string())
        return {"imported": 0, "skipped": 0, "errors": 1}

    logger.info("Found header row at Excel row %d", header_row + 1)

    try:
        df = pd.read_excel(filepath, sheet_name=SHEET_NAME, header=header_row)
    except Exception as e:
        logger.error("Could not re-read with header=%d: %s", header_row, e)
        return {"imported": 0, "skipped": 0, "errors": 1}

    logger.info("Rows read: %d   Columns: %s", len(df), list(df.columns))

    # Rename columns to standard internal names
    col_rename = {
        "Date1":      "entry_date",
        "Name":       "symbol",
        "Status":     "status",
        "QTY":        "qty",
        "BUY":        "entry_price",
        "Value":      "value",
        "SL":         "stop_loss",
        "POS":        "direction",
        "Date2":      "exit_date",
        "SELL":       "exit_price",
        "Gain/Loss":  "pl_pkr",
        "Gain/Loss%": "pl_pct",
        "Days":       "holding_days",
        "R:R":        "actual_rr",
    }
    df = df.rename(columns={k: v for k, v in col_rename.items() if k in df.columns})

    imported = skipped = errors = not_initiated = updated = 0

    with get_conn() as conn:
        for idx, row in df.iterrows():
            try:
                # Skip rows with no symbol
                symbol = str(row.get("symbol", "")).strip().upper()
                if not symbol or symbol in ("NAN", "NAME", ""):
                    continue

                # Skip not-initiated trades (Status blank)
                status_raw = str(row.get("status", "")).strip().upper()
                if status_raw not in ("C", "O"):
                    not_initiated += 1
                    continue

                entry_dt = _parse_date(row.get("entry_date"))
                if not entry_dt:
                    logger.warning("Row %d (%s): no entry date — skipped", idx + 2, symbol)
                    errors += 1
                    continue

                direction    = _normalise_direction(row.get("direction", "Long"))
                entry_price  = float(row["entry_price"]) if pd.notna(row.get("entry_price")) else None
                exit_price   = float(row["exit_price"])  if pd.notna(row.get("exit_price"))  else None
                exit_dt      = _parse_date(row.get("exit_date"))
                stop_loss    = float(row["stop_loss"])   if pd.notna(row.get("stop_loss"))   else None
                qty          = float(row["qty"])          if pd.notna(row.get("qty"))          else None
                holding_days = int(row["holding_days"])  if pd.notna(row.get("holding_days")) else None
                actual_rr    = float(row["actual_rr"])   if pd.notna(row.get("actual_rr"))    else None
                pl_pkr       = float(row["pl_pkr"])      if pd.notna(row.get("pl_pkr"))       else None

                # P&L % — prefer column value, calculate as fallback
                pl_pct = None
                if pd.notna(row.get("pl_pct")):
                    raw = row["pl_pct"]
                    pl_pct = float(raw) * 100 if abs(float(raw)) < 5 else float(raw)
                elif entry_price and exit_price and entry_price > 0:
                    pl_pct = (
                        (exit_price - entry_price) / entry_price * 100
                        if direction == "LONG"
                        else (entry_price - exit_price) / entry_price * 100
                    )

                outcome = _outcome_from_pl(pl_pct) if status_raw == "C" else None
                db_status = "Closed" if status_raw == "C" else "Active"

                ep = entry_price or 0.0
                sl = stop_loss or ep * 0.94
                risk_pct = abs((ep - sl) / ep * 100) if ep > 0 else 0.0

                # Check for existing record
                existing = conn.execute(
                    """SELECT id, status, outcome, actual_pl_pkr, actual_pl_pct, actual_rr
                       FROM trade_setups
                       WHERE symbol=? AND created_date=? AND direction=? AND source='Actual'""",
                    (symbol, entry_dt, direction),
                ).fetchone()

                if existing:
                    # UPDATE if status, outcome, or P&L values differ from Excel
                    existing_status  = existing["status"]  if existing["status"]  else ""
                    existing_outcome = existing["outcome"] if existing["outcome"] else ""
                    existing_pl_pkr  = existing["actual_pl_pkr"] or 0
                    existing_pl_pct  = existing["actual_pl_pct"] or 0

                    pl_pkr_drift = abs(existing_pl_pkr - (pl_pkr or 0)) > 0.5
                    pl_pct_drift = abs(existing_pl_pct - (pl_pct or 0)) > 0.01

                    needs_update = (
                        existing_status  != db_status or
                        existing_outcome != (outcome or "") or
                        pl_pkr_drift or
                        pl_pct_drift
                    )

                    if not needs_update:
                        skipped += 1
                        continue

                    if dry_run:
                        logger.info(
                            "  [DRY RUN UPDATE] %s %s | %s→%s | outcome %s→%s",
                            direction, symbol,
                            existing_status, db_status,
                            existing_outcome or "none", outcome or "none",
                        )
                        imported += 1
                        continue

                    conn.execute(
                        """UPDATE trade_setups SET
                               status=?, outcome=?,
                               actual_exit=?, actual_pl_pct=?, actual_pl_pkr=?,
                               holding_days=?, actual_rr=?, exit_date=?,
                               entry_price=?, actual_entry=?, stop_loss=?, risk_pct=?
                           WHERE id=?""",
                        (
                            db_status, outcome,
                            exit_price, pl_pct, pl_pkr,
                            holding_days, actual_rr, exit_dt,
                            ep, ep, sl, risk_pct,
                            existing["id"],
                        ),
                    )
                    imported += 1
                    logger.info(
                        "  UPDATED %s %s | %s→%s | outcome: %s | P&L: %s%%",
                        direction, symbol,
                        existing_status, db_status,
                        outcome or "open",
                        f"{pl_pct:+.1f}" if pl_pct is not None else "?",
                    )
                    continue

                # New record — INSERT
                if dry_run:
                    pl_str = f"{pl_pct:+.1f}%" if pl_pct is not None else "?"
                    logger.info(
                        "  [DRY RUN INSERT] %s %s | entry %s | exit %s | P&L %s | %s | %s",
                        direction, symbol, entry_dt, exit_dt or "open",
                        pl_str, outcome or "open", db_status
                    )
                    imported += 1
                    continue

                conn.execute(
                    """INSERT INTO trade_setups
                        (created_date, direction, symbol, sector, sector_momentum,
                         stock_perf_30d, stock_perf_10d, latest_close,
                         entry_price, stop_loss, target_1r, target_2r,
                         risk_pct, atr_pct,
                         status, outcome,
                         actual_entry, actual_exit, actual_pl_pct, actual_pl_pkr,
                         holding_days, actual_rr,
                         exit_date, source)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        entry_dt, direction, symbol,
                        "Unknown", "—",          # sector/momentum not in Excel
                        0.0, 0.0, ep,
                        ep, sl,
                        ep * 1.07, ep * 1.14,    # 7% and 14% targets from your Excel headers
                        risk_pct, 0.0,
                        db_status, outcome,
                        ep, exit_price, pl_pct, pl_pkr,
                        holding_days, actual_rr,
                        exit_dt, "Actual",
                    ),
                )
                imported += 1

            except Exception as e:
                logger.warning("Row %d: error — %s", idx + 2, e)
                errors += 1

    # ── Fix existing Breakeven records that should be Win or Loss ────────────────
    # Previous threshold was ±0.25% — reclassify anything non-zero
    if not dry_run:
        with get_conn() as conn:
            fixed = conn.execute(
                """UPDATE trade_setups
                   SET outcome = CASE
                       WHEN actual_pl_pct > 0 THEN 'Win'
                       WHEN actual_pl_pct < 0 THEN 'Loss'
                       ELSE 'Breakeven'
                   END
                   WHERE source = 'Actual'
                     AND outcome = 'Breakeven'
                     AND actual_pl_pct IS NOT NULL
                     AND actual_pl_pct != 0"""
            ).rowcount
            if fixed:
                logger.info("  Reclassified %d Breakeven → Win/Loss", fixed)

    # ── Sync portfolio value from JOURNAL-2!B3 → portfolio_values ────────────────
    # B3 holds the current portfolio value, manually updated daily in Excel.
    # Writing it here keeps the DB current so Streamlit Cloud has a fresh capital
    # base for Kelly sizing without any manual Audit-page entry.
    if not dry_run:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
            ws = wb[SHEET_NAME]
            raw_val = ws["B3"].value
            wb.close()
            if raw_val is not None:
                portfolio_val = float(raw_val)
                if portfolio_val > 0:
                    from database import add_portfolio_value
                    today = pd.Timestamp.today().strftime("%Y-%m-%d")
                    add_portfolio_value(today, portfolio_val, notes="Auto-synced from JOURNAL-2!B3")
                    logger.info("Portfolio value synced: PKR %.0f (as of %s)", portfolio_val, today)
                else:
                    logger.warning("B3 value is zero or negative (%.0f) — portfolio_values not updated", portfolio_val)
            else:
                logger.warning("B3 is empty — portfolio_values not updated")
        except Exception as e:
            logger.warning("Could not sync portfolio value from B3: %s", e)

    logger.info(
        "\n=== SUMMARY ===\n"
        "  Imported   : %d\n"
        "  Skipped    : %d  (no change needed)\n"
        "  Not taken  : %d  (Status blank — trades never initiated, skipped)\n"
        "  Errors     : %d",
        imported, skipped, not_initiated, errors
    )
    if dry_run:
        logger.info("\n  *** DRY RUN — nothing written to DB. Remove --dry-run to import. ***")

    return {"imported": imported, "skipped": skipped, "errors": errors}


def main():
    parser = argparse.ArgumentParser(
        description="Import Journal-2 trades from ASSET ALLOCATION.xlsx into psx_data.db"
    )
    parser.add_argument("--file", required=True, help="Full path to ASSET ALLOCATION.xlsx")
    parser.add_argument("--dry-run", action="store_true", help="Preview without writing to DB")
    args = parser.parse_args()

    if not os.path.exists(args.file):
        logger.error("File not found: %s", args.file)
        sys.exit(1)

    init_db()
    import_trades(args.file, dry_run=args.dry_run)


if __name__ == "__main__":
    main()
